import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_out = r"c:\Users\USER\Documents\PUSDATIN\Laporan_Anomali_Siasati_2026.xlsx"
print(f"Menyiapkan pembaruan file Excel: {excel_out}...")

# 1. Baca data dari CSV
print("Membaca CSV...")
df_exact = pd.read_csv(r"c:\Users\USER\Documents\PUSDATIN\data_duplikat_sama_persis.csv", low_memory=False)
df_diff = pd.read_csv(r"c:\Users\USER\Documents\PUSDATIN\data_duplikat_biasa_beda_nilai.csv", low_memory=False)
df_empty_daily = pd.read_csv(r"c:\Users\USER\Documents\PUSDATIN\data_lat_lon_kosong.csv", low_memory=False)

# Buat ExcelWriter dengan engine openpyxl
with pd.ExcelWriter(excel_out, engine='openpyxl') as writer:
    # Buat placeholder sheet summary terlebih dahulu
    df_summary_placeholder = pd.DataFrame({'Informasi': ['Memuat ringkasan...']})
    df_summary_placeholder.to_excel(writer, sheet_name='Ringkasan Eksekutif', index=False)
    
    print("Menulis sheet 1: Duplikat_Sama_Persis (7,226 baris)...")
    df_exact.to_excel(writer, sheet_name='Duplikat_Sama_Persis', index=False)
    
    print("Menulis sheet 2: Duplikat_Beda_Nilai (34,229 baris)...")
    df_diff.to_excel(writer, sheet_name='Duplikat_Beda_Nilai', index=False)
    
    print("Menulis sheet 3: Lat_Long_Kosong (8,103 baris)...")
    df_empty_daily.to_excel(writer, sheet_name='Lat_Long_Kosong', index=False)

print("Memuat workbook untuk styling & desain Ringkasan Eksekutif...")
wb = openpyxl.load_workbook(excel_out)

# -------------------------------------------------------------
# STYLING SHEET DATA
# -------------------------------------------------------------
sheet_configs = {
    'Duplikat_Sama_Persis': {'header_color': '1F4E79'},   # Navy Blue
    'Duplikat_Beda_Nilai': {'header_color': '2E75B6'},    # Medium Blue
    'Lat_Long_Kosong': {'header_color': 'C00000'},        # Crimson Red
}

header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=10)
thin_border = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

for sname, cfg in sheet_configs.items():
    if sname in wb.sheetnames:
        ws = wb[sname]
        ws.freeze_panes = 'A2'
        fill_color = PatternFill(start_color=cfg['header_color'], end_color=cfg['header_color'], fill_type='solid')
        
        # Style header row
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = fill_color
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 28
        
        # Enable auto filter
        ws.auto_filter.ref = ws.dimensions
        
        # Auto-adjust column width (scan first 100 rows)
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for idx, cell in enumerate(col[:100]):
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -------------------------------------------------------------
# DESAIN SHEET 'Ringkasan Eksekutif'
# -------------------------------------------------------------
ws_sum = wb['Ringkasan Eksekutif']
# Bersihkan cell
for row in ws_sum['A1:Z60']:
    for cell in row:
        cell.value = None

ws_sum.views.sheetView[0].showGridLines = True

# Palet warna
c_navy = '0F2537'
c_blue = '1F4E79'
c_accent = '2E75B6'
c_light = 'F2F4F7'
c_gray = '595959'

font_title = Font(name='Segoe UI', size=16, bold=True, color=c_navy)
font_subtitle = Font(name='Segoe UI', size=11, italic=True, color=c_gray)
font_sec_hdr = Font(name='Segoe UI', size=12, bold=True, color='FFFFFF')
font_tbl_hdr = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_tbl_bold = Font(name='Segoe UI', size=10, bold=True)
font_tbl_reg = Font(name='Segoe UI', size=10)

fill_sec = PatternFill(start_color=c_blue, end_color=c_blue, fill_type='solid')
fill_th = PatternFill(start_color=c_accent, end_color=c_accent, fill_type='solid')
fill_alt = PatternFill(start_color=c_light, end_color=c_light, fill_type='solid')

# Judul
ws_sum['B2'] = "LAPORAN AUDIT ANOMALI DATA SIASATI KEMENHUB 2026"
ws_sum['B2'].font = font_title
ws_sum['B3'] = "Cakupan Data: 1 Januari 2026 s.d. 28 September 2026 (Total 302.089 Baris Data Mentah)"
ws_sum['B3'].font = font_subtitle

# 1. Ringkasan Eksekutif Metrik Utama
ws_sum['B5'] = "1. RINGKASAN TEMUAN UTAMA"
ws_sum['B5'].font = font_sec_hdr
ws_sum['B5'].fill = fill_sec
for c in ['C5', 'D5', 'E5', 'F5', 'G5']:
    ws_sum[c].fill = fill_sec
ws_sum.row_dimensions[5].height = 24

headers_kpi = ["No", "Kategori Anomali / Temuan", "Jumlah Baris", "Satuan Grup / Unik", "Sheet Excel", "Status Kritis"]
for i, h in enumerate(headers_kpi, start=2):
    cell = ws_sum.cell(row=6, column=i, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_th
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws_sum.row_dimensions[6].height = 22

kpi_data = [
    [1, "Duplikat Sama Persis (Kembar Identik)", "7.226 baris", "3.613 pasang", "Duplikat_Sama_Persis", "Tinggi (Double Entry)"],
    [2, "- Bernilai Positif Nyata (> 0)", "160 baris", "80 pasang", "Duplikat_Sama_Persis", "Sangat Tinggi (Merusak Agregasi)"],
    [3, "- Bernilai Nol Semua (0 + 0)", "7.066 baris", "3.533 pasang", "Duplikat_Sama_Persis", "Sedang (Sampah Dummy)"],
    [4, "Duplikat Beda Nilai (1 Riil + Dummy Nol)", "34.229 baris", "16.099 tanggal-simpul", "Duplikat_Beda_Nilai", "Tinggi (Integritas Data)"],
    [5, "- Baris Riil Transaksi Aktif", "17.243 baris", "-", "Duplikat_Beda_Nilai", "Perlu Dipelihara"],
    [6, "- Baris Dummy Nol Sistem", "16.986 baris", "-", "Duplikat_Beda_Nilai", "Perlu Dieliminasi"],
    [7, "Data Lat Long Kosong (Blank/NaN)", "8.103 baris", "177 prasarana unik", "Lat_Long_Kosong", "Sedang (Aset Spasial)"],
]

for row_idx, row_vals in enumerate(kpi_data, start=7):
    is_alt = (row_idx % 2 == 0)
    for col_idx, val in enumerate(row_vals, start=2):
        cell = ws_sum.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_tbl_bold if row_idx in [7, 10, 13] else font_tbl_reg
        if is_alt:
            cell.fill = fill_alt
        cell.border = thin_border
        if col_idx in [2, 4, 5, 6, 7]:
            cell.alignment = Alignment(horizontal='center' if col_idx in [2, 4, 6, 7] else 'left', vertical='center')

# 2. Distribusi per Moda Transportasi
ws_sum['B16'] = "2. DISTRIBUSI TEMUAN PER MODA TRANSPORTASI"
ws_sum['B16'].font = font_sec_hdr
ws_sum['B16'].fill = fill_sec
for c in ['C16', 'D16', 'E16', 'F16', 'G16']:
    ws_sum[c].fill = fill_sec
ws_sum.row_dimensions[16].height = 24

headers_moda = ["Moda Transportasi", "Duplikat Sama Persis (Baris)", "Duplikat Beda Nilai (Baris)", "Lat Long Kosong (Harian)", "Simpul Unik Tanpa Lat Long", "Catatan Kritis"]
for i, h in enumerate(headers_moda, start=2):
    cell = ws_sum.cell(row=17, column=i, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_th
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws_sum.row_dimensions[17].height = 22

moda_breakdown = [
    ["ASDP (Penyeberangan)", "0 baris", "12 baris", "4.204 baris", "77 simpul", "Dominan masalah koordinat kosong (Ulee Lheu, Ciwandan, Bau Bau)"],
    ["BUS (Terminal)", "66 baris", "0 baris", "3.263 baris", "63 simpul", "66 baris duplikat nyata di Terminal Bangsri & Pecangaan"],
    ["KA (Kereta Api)", "7.158 baris", "32.089 baris", "103 baris", "6 simpul", "Paling masif duplikat 1 riil + 1 dummy nol (DJKA / KAI)"],
    ["LAUT (Pelabuhan)", "2 baris", "2.128 baris", "1 baris", "1 simpul", "2.128 baris pencatatan ganda beda rute/layanan (domestik vs intl)"],
    ["UDARA (Bandara)", "0 baris", "0 baris", "532 baris", "30 simpul", "0 duplikat; 30 bandara perintis tanpa koordinat lat/lon"],
    ["TOTAL MULTIMODA", "7.226 baris", "34.229 baris", "8.103 baris", "177 simpul", "Data per 28 September 2026"],
]

for row_idx, row_vals in enumerate(moda_breakdown, start=18):
    is_last = (row_idx == 23)
    for col_idx, val in enumerate(row_vals, start=2):
        cell = ws_sum.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_tbl_bold if is_last else font_tbl_reg
        cell.border = thin_border
        if is_last:
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        elif row_idx % 2 == 1:
            cell.fill = fill_alt
        if col_idx in [3, 4, 5, 6]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

# 3. Panduan Penggunaan Sheet
ws_sum['B25'] = "3. PANDUAN PENGGUNAAN SHEET DALAM WORKBOOK INI"
ws_sum['B25'].font = font_sec_hdr
ws_sum['B25'].fill = fill_sec
for c in ['C25', 'D25', 'E25', 'F25', 'G25']:
    ws_sum[c].fill = fill_sec
ws_sum.row_dimensions[25].height = 24

headers_guide = ["Nama Sheet", "Deskripsi Konten", "Cara Analisis & Rekomendasi Tindakan"]
for i, h in enumerate(headers_guide, start=2):
    cell = ws_sum.cell(row=26, column=i, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_th
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws_sum.row_dimensions[26].height = 22

guides = [
    ["Duplikat_Sama_Persis", "Berisi 7.226 baris di mana 2 baris rekaman memiliki tanggal, stasiun, dan angka metrik yang sama persis.", "Filter kolom 'kategori_duplikat'. Hapus 1 dari 2 baris kembar agar tidak terjadi overcounting / inflasi statistik penumpang."],
    ["Duplikat_Beda_Nilai", "Berisi 34.229 baris di mana pada 1 simpul dan tanggal ada 2 baris pencatatan dengan nilai berbeda.", "Lihat kolom 'peran_baris': Baris Dummy Nol Sistem (16.986 baris) dapat di-drop/dibersihkan. Baris Riil Transaksi (17.243 baris) dipertahankan."],
    ["Lat_Long_Kosong", "Berisi 8.103 baris transaksi harian simpul yang belum memiliki koordinat lintang & bujur (mencakup 177 simpul unik).", "Gunakan untuk identifikasi transaksi yang belum terpetakan di GIS/Peta spasial. Bisa langsung di-pivot/remove duplicates di Excel jika ingin melihat daftar prasarananya."]
]

for row_idx, row_vals in enumerate(guides, start=27):
    for col_idx, val in enumerate(row_vals, start=2):
        cell = ws_sum.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_tbl_bold if col_idx == 2 else font_tbl_reg
        cell.border = thin_border
        if row_idx % 2 == 0:
            cell.fill = fill_alt
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws_sum.row_dimensions[row_idx].height = 36

# Lebar kolom ringkasan eksekutif
col_widths = {
    'A': 4,
    'B': 8,
    'C': 34,
    'D': 22,
    'E': 26,
    'F': 28,
    'G': 40
}
for c, w in col_widths.items():
    ws_sum.column_dimensions[c].width = w

# Simpan workbook
wb.save(excel_out)
print(f"Workbook berhasil disimpan (tanpa sheet simpul unik): {excel_out}")
